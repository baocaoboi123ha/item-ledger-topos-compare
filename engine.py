# -*- coding: utf-8 -*-
"""So sánh Item Ledger Entries vs Topos theo Ngày + Location + Type + Mã Nội Bộ."""

from __future__ import annotations

import re
import shutil
import unicodedata
from collections import defaultdict
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter

ENGINE_VERSION = "1.1"
QTY_TOLERANCE = 0.001
OUTPUT_SUFFIX = " - ILE vs Topos.xlsx"
RESULT_SHEET = "Ket qua lech"
SUMMARY_SHEET = "TongHop"
WEB_PREVIEW_ROWS = 300

RESULT_HEADERS = (
    "Ngày",
    "Location",
    "Type",
    "Item No/Mã Nội Bộ",
    "Quantity Item Ledger",
    "Quantity Topos",
    "Lệch (Topos - Item Ledger)",
    "Status",
)
SUMMARY_HEADERS = (
    "Ngày",
    "Location",
    "Type",
    "Số key lệch",
    "Tổng Quantity Item Ledger",
    "Tổng Quantity Topos",
    "Tổng lệch (Topos - ILE)",
)

ILE_SHEET_ALIASES = (
    "item ledger entries",
    "itemledgerentries",
    "item ledger",
    "ile",
    "ledger",
)
TOPOS_SHEET_ALIASES = ("topos",)

ILE_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("postingdate", "ngay", "date", "ngayhachtoan"),
    "location": ("locationcode", "location", "tencuahang", "bustorecode"),
    "type": ("type",),
    "item": ("itemno", "manoibo", "masanpham"),
    "qty": ("quantity", "soluong", "qty"),
}
TOPOS_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("ngay", "date", "postingdate"),
    "location": ("tencuahang", "locationcode", "location"),
    "type": ("type",),
    "item": ("manoibo", "itemno", "masanpham"),
    "qty": ("soluong", "quantity", "qty"),
}
REQUIRED_ROLES = ("date", "location", "type", "item", "qty")

_THIN = Side(style="thin", color="808080")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F5179")
_MISS_ILE = PatternFill("solid", fgColor="FCE5D6")
_MISS_TP = PatternFill("solid", fgColor="D6EAF8")
_DIFF_FILL = PatternFill("solid", fgColor="F8D4B8")
QTY_FMT = "#,##0.###"
LogFn = Callable[[str], None]


class CompareError(Exception):
    """Lỗi dữ liệu / thiếu cột để người dùng đọc được."""


def normalize_header(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = text.replace("\u0111", "d").replace("\u0110", "d")
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "", text)


def normalize_sheet(name: str) -> str:
    return re.sub(r"\s+", " ", normalize_header(name))


def parse_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if 20000 < number < 90000:
            try:
                return (datetime(1899, 12, 30) + timedelta(days=number)).date()
            except (OverflowError, ValueError):
                pass
    text = str(value).strip()
    if not text or text.upper() in ("NULL", "NONE", "NAN"):
        return None
    if "T" in text:
        text = text.split("T", 1)[0]
    if " " in text and re.match(r"^\d", text):
        text = text.split(" ", 1)[0]
    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d.%m.%Y",
        "%m/%d/%Y",
        "%Y/%m/%d",
        "%d/%m/%y",
        "%d.%m.%y",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    m = re.match(r"^(\d{4})(\d{2})(\d{2})$", text)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def to_float(value) -> float:
    if value is None or value == "" or isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        cleaned = "".join(ch for ch in text if ch.isdigit() or ch in ".-")
        if cleaned not in {"", "-", "."}:
            try:
                return float(cleaned)
            except ValueError:
                return 0.0
        return 0.0


def key_text(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        text = str(value).strip()
        return text[:-2] if text.endswith(".0") else text
    return str(value).strip()


def resolve_columns(
    headers: Iterable,
    alias_map: dict[str, tuple[str, ...]],
    source_label: str,
) -> dict[str, int]:
    norm_index: dict[str, int] = {}
    for i, header in enumerate(headers):
        key = normalize_header(header)
        if key and key not in norm_index:
            norm_index[key] = i
    resolved: dict[str, int] = {}
    for role, keys in alias_map.items():
        for key in keys:
            if key in norm_index:
                resolved[role] = norm_index[key]
                break
    missing = [role for role in REQUIRED_ROLES if role not in resolved]
    if missing:
        shown = ", ".join(str(h) for h in headers if h is not None)
        raise CompareError(
            f"{source_label}: thiếu cột {', '.join(missing)}. Header có: {shown}"
        )
    return resolved


def _find_sheet(wb, aliases: tuple[str, ...], label: str, fallback_single: bool) -> str:
    by_norm = {normalize_sheet(name): name for name in wb.sheetnames}
    for alias in aliases:
        hit = by_norm.get(normalize_sheet(alias))
        if hit:
            return hit
    for name in wb.sheetnames:
        n = normalize_sheet(name)
        if any(normalize_sheet(alias) in n or n in normalize_sheet(alias) for alias in aliases):
            return name
    if fallback_single and len(wb.sheetnames) == 1:
        return wb.sheetnames[0]
    raise CompareError(
        f"Không tìm thấy sheet {label}. Có: {', '.join(wb.sheetnames)}"
    )


def _read_sheet(wb, sheet_name: str) -> list[tuple]:
    return list(wb[sheet_name].iter_rows(values_only=True))


def _pivot(
    rows: list[tuple],
    cols: dict[str, int],
    label: str,
    log: LogFn,
) -> dict[tuple, float]:
    if not rows:
        raise CompareError(f"{label}: không có dòng dữ liệu.")
    pivot: dict[tuple, float] = defaultdict(float)
    skipped_date = 0
    skipped_item = 0
    used = 0
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        d = parse_date(row[cols["date"]])
        if d is None:
            skipped_date += 1
            continue
        item = key_text(row[cols["item"]])
        if not item:
            skipped_item += 1
            continue
        loc = key_text(row[cols["location"]])
        typ = key_text(row[cols["type"]])
        qty = to_float(row[cols["qty"]])
        pivot[(d, loc, typ, item)] += qty
        used += 1
    if skipped_date:
        log(f"{label}: bỏ qua {skipped_date} dòng không đọc được ngày.")
    if skipped_item:
        log(f"{label}: bỏ qua {skipped_item} dòng thiếu Item No/Mã Nội Bộ.")
    log(f"{label}: {used} dòng → {len(pivot)} key.")
    return pivot


def _qty_mismatch(a: float, b: float) -> bool:
    return abs(b - a) > QTY_TOLERANCE


def _status(ile_qty: float, tp_qty: float, in_ile: bool, in_tp: bool) -> str:
    if in_ile and not in_tp:
        return "Chỉ có Item Ledger"
    if in_tp and not in_ile:
        return "Chỉ có Topos"
    return "Lệch số lượng"


def _load_pivots(
    ile_path: Path,
    topos_path: Path | None,
    log: LogFn,
) -> tuple[dict[tuple, float], dict[tuple, float]]:
    ile_wb = load_workbook(ile_path, read_only=True, data_only=True)
    try:
        same_file = topos_path is None or Path(topos_path).resolve() == Path(ile_path).resolve()
        ile_sheet = _find_sheet(ile_wb, ILE_SHEET_ALIASES, "Item Ledger Entries", fallback_single=not same_file)
        ile_rows = _read_sheet(ile_wb, ile_sheet)
        log(f"Item Ledger: sheet '{ile_sheet}' ({ile_path.name})")
        ile_cols = resolve_columns(ile_rows[0] if ile_rows else (), ILE_ALIASES, "Item Ledger")
        ile_pivot = _pivot(ile_rows, ile_cols, "Item Ledger", log)

        if same_file:
            topos_sheet = _find_sheet(ile_wb, TOPOS_SHEET_ALIASES, "Topos", fallback_single=False)
            topos_rows = _read_sheet(ile_wb, topos_sheet)
            log(f"Topos: sheet '{topos_sheet}' (cùng file)")
        else:
            topos_wb = load_workbook(topos_path, read_only=True, data_only=True)
            try:
                topos_sheet = _find_sheet(
                    topos_wb, TOPOS_SHEET_ALIASES, "Topos", fallback_single=True
                )
                topos_rows = _read_sheet(topos_wb, topos_sheet)
                log(f"Topos: sheet '{topos_sheet}' ({Path(topos_path).name})")
            finally:
                topos_wb.close()
        topos_cols = resolve_columns(topos_rows[0] if topos_rows else (), TOPOS_ALIASES, "Topos")
        topos_pivot = _pivot(topos_rows, topos_cols, "Topos", log)
    finally:
        ile_wb.close()
    return ile_pivot, topos_pivot


def build_mismatch_rows(
    ile_pivot: dict[tuple, float],
    topos_pivot: dict[tuple, float],
) -> list[tuple]:
    keys = set(ile_pivot) | set(topos_pivot)
    rows: list[tuple] = []
    for key in sorted(keys, key=lambda k: (k[0] or date.min, k[1], k[2], k[3])):
        in_ile = key in ile_pivot
        in_tp = key in topos_pivot
        ile_qty = ile_pivot.get(key, 0.0)
        tp_qty = topos_pivot.get(key, 0.0)
        if not _qty_mismatch(ile_qty, tp_qty):
            continue
        d, loc, typ, item = key
        rows.append(
            (
                d,
                loc,
                typ,
                item,
                ile_qty,
                tp_qty,
                tp_qty - ile_qty,
                _status(ile_qty, tp_qty, in_ile, in_tp),
            )
        )
    return rows


def _summary_rows(mismatch: list[tuple]) -> list[tuple]:
    grouped: dict[tuple, list[tuple]] = defaultdict(list)
    for row in mismatch:
        grouped[(row[0], row[1], row[2])].append(row)
    out = []
    for key in sorted(grouped, key=lambda k: (k[0] or date.min, k[1], k[2])):
        items = grouped[key]
        ile_sum = sum(r[4] for r in items)
        tp_sum = sum(r[5] for r in items)
        out.append((key[0], key[1], key[2], len(items), ile_sum, tp_sum, tp_sum - ile_sum))
    return out


def _style_header(ws, headers: tuple[str, ...]) -> None:
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = _BORDER


def _write_table(ws, headers: tuple[str, ...], rows: list[tuple], qty_cols: set[int]) -> None:
    _style_header(ws, headers)
    for r_idx, row in enumerate(rows, 2):
        status = row[7] if len(row) > 7 else ""
        fill = None
        if status == "Chỉ có Item Ledger":
            fill = _MISS_TP
        elif status == "Chỉ có Topos":
            fill = _MISS_ILE
        elif status == "Lệch số lượng":
            fill = _DIFF_FILL
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = _BORDER
            if fill is not None:
                cell.fill = fill
            if c_idx in qty_cols and isinstance(value, (int, float)):
                cell.number_format = QTY_FMT
            if c_idx == 1 and isinstance(value, date):
                cell.number_format = "YYYY-MM-DD"
    widths = [14, 14, 10, 22, 22, 18, 26, 22]
    for i, width in enumerate(widths[: len(headers)], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    last_col = get_column_letter(len(headers))
    last_row = max(1, len(rows) + 1)
    ws.auto_filter = AutoFilter(ref=f"A1:{last_col}{last_row}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"


def _unique_sheet_name(name: str, used: set[str]) -> str:
    if name not in used:
        return name
    for i in range(2, 100):
        candidate = f"{name[:28]} ({i})"
        if candidate not in used:
            return candidate
    raise CompareError(f"Không đặt được tên sheet cho '{name}'.")


def _copy_worksheet(source_ws, target_ws) -> None:
    for row in source_ws.iter_rows():
        for cell in row:
            new_cell = target_ws.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.border = copy(cell.border)
                new_cell.fill = copy(cell.fill)
                new_cell.number_format = cell.number_format
                new_cell.protection = copy(cell.protection)
                new_cell.alignment = copy(cell.alignment)
    for merged in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged))
    for letter, dim in source_ws.column_dimensions.items():
        if dim.width:
            target_ws.column_dimensions[letter].width = dim.width
    for idx, dim in source_ws.row_dimensions.items():
        if dim.height:
            target_ws.row_dimensions[idx].height = dim.height
    target_ws.freeze_panes = source_ws.freeze_panes
    if source_ws.auto_filter and source_ws.auto_filter.ref:
        target_ws.auto_filter.ref = source_ws.auto_filter.ref
    target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines


def _strip_result_sheets(wb) -> None:
    for name in list(wb.sheetnames):
        if name in (RESULT_SHEET, SUMMARY_SHEET):
            del wb[name]


def _append_workbook_sheets(src_path: Path, dest_wb, used: set[str], log: LogFn) -> None:
    src = load_workbook(src_path)
    try:
        for ws in src.worksheets:
            if ws.title in (RESULT_SHEET, SUMMARY_SHEET):
                continue
            title = _unique_sheet_name(ws.title, used)
            new_ws = dest_wb.create_sheet(title)
            _copy_worksheet(ws, new_ws)
            used.add(title)
            if title != ws.title:
                log(f"Đổi tên sheet '{ws.title}' → '{title}' để tránh trùng.")
    finally:
        src.close()


def write_result_workbook(
    output: Path,
    mismatch: list[tuple],
    log: LogFn,
    source: Path,
    extra: Path | None = None,
) -> None:
    """File mới = toàn bộ sheet gốc (trước) + sheet kết quả (sau)."""
    output.parent.mkdir(parents=True, exist_ok=True)
    same_file = extra is None or Path(extra).resolve() == Path(source).resolve()

    shutil.copy2(source, output)
    wb = load_workbook(output)
    try:
        _strip_result_sheets(wb)
        used = set(wb.sheetnames)
        if not same_file:
            log(f"Ghép thêm sheet từ file Topos: {extra.name}")
            _append_workbook_sheets(extra, wb, used, log)

        ws = wb.create_sheet(RESULT_SHEET)
        _write_table(ws, RESULT_HEADERS, mismatch, {5, 6, 7})
        ws2 = wb.create_sheet(SUMMARY_SHEET)
        _write_table(ws2, SUMMARY_HEADERS, _summary_rows(mismatch), {5, 6, 7})

        order = [n for n in wb.sheetnames if n not in (RESULT_SHEET, SUMMARY_SHEET)]
        order.extend([RESULT_SHEET, SUMMARY_SHEET])
        for idx, name in enumerate(order):
            wb.move_sheet(name, offset=idx - wb.sheetnames.index(name))

        wb.save(output)
        log(
            f"Ghi file gộp: {output.name} — sheet gốc: {', '.join(order[:-2])}; "
            f"sau đó {RESULT_SHEET}, {SUMMARY_SHEET} ({len(mismatch)} dòng lệch)."
        )
    finally:
        wb.close()


def compare_workbook(
    ile_path,
    output_path=None,
    log: LogFn | None = None,
    topos_path=None,
) -> dict:
    def _log(msg: str) -> None:
        if log:
            log(str(msg))

    source = Path(ile_path)
    if not source.exists():
        raise CompareError(f"Không thấy file Item Ledger: {source}")
    extra = Path(topos_path) if topos_path else None
    if extra is not None and not extra.exists():
        raise CompareError(f"Không thấy file Topos: {extra}")

    dest = Path(output_path) if output_path else source.with_name(source.stem + OUTPUT_SUFFIX)
    _log(f"Engine {ENGINE_VERSION}")
    _log("Key: Ngày + Location/Tên Cửa Hàng + Type + Item No/Mã Nội Bộ")
    _log("Lệch = Quantity Topos − Quantity Item Ledger")

    ile_pivot, topos_pivot = _load_pivots(source, extra, _log)
    mismatch = build_mismatch_rows(ile_pivot, topos_pivot)
    only_ile = sum(1 for r in mismatch if r[7] == "Chỉ có Item Ledger")
    only_tp = sum(1 for r in mismatch if r[7] == "Chỉ có Topos")
    both = sum(1 for r in mismatch if r[7] == "Lệch số lượng")
    _log(
        f"Key ILE {len(ile_pivot)}, Topos {len(topos_pivot)}, "
        f"lệch {len(mismatch)} (cả hai {both}, chỉ ILE {only_ile}, chỉ Topos {only_tp})."
    )
    write_result_workbook(dest, mismatch, _log, source, extra)

    preview = []
    for row in mismatch[:WEB_PREVIEW_ROWS]:
        preview.append(
            {
                "Ngày": row[0].isoformat() if isinstance(row[0], date) else row[0],
                "Location": row[1],
                "Type": row[2],
                "Item No/Mã Nội Bộ": row[3],
                "Quantity Item Ledger": row[4],
                "Quantity Topos": row[5],
                "Lệch (Topos - Item Ledger)": round(row[6], 6),
                "Status": row[7],
            }
        )
    return {
        "OutputPath": str(dest),
        "EngineVersion": ENGINE_VERSION,
        "IleKeys": len(ile_pivot),
        "ToposKeys": len(topos_pivot),
        "MismatchRows": len(mismatch),
        "OnlyIle": only_ile,
        "OnlyTopos": only_tp,
        "QtyDiffBoth": both,
        "PreviewRows": preview,
        "ResultHeaders": list(RESULT_HEADERS),
    }
